import io
import os
import zipfile
from pypdf import PdfReader
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
from docx import Document
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from utils import get_temp_path


def pdf_to_images(input_path: str, fmt: str = "png", dpi: int = 150) -> str:
    reader = PdfReader(input_path)
    output_dir = get_temp_path()
    os.makedirs(output_dir)

    image_paths = []
    for i, page in enumerate(reader.pages):
        for img_obj in page.images:
            try:
                pil_img = Image.open(io.BytesIO(img_obj.data))
                img_path = os.path.join(output_dir, f"page_{i + 1}.{fmt}")
                if fmt == "jpg":
                    if pil_img.mode in ("RGBA", "P"):
                        pil_img = pil_img.convert("RGB")
                    pil_img.save(img_path, "JPEG", quality=90)
                else:
                    pil_img.save(img_path, "PNG")
                image_paths.append(img_path)
                break
            except Exception:
                continue

    if not image_paths:
        # Fallback: render pages using pypdf's basic extraction
        # For better rendering, we create a simple representation
        from reportlab.pdfgen import canvas as rl_canvas

        for i, page in enumerate(reader.pages):
            text = page.extract_text() or f"Page {i + 1}"
            img_path = os.path.join(output_dir, f"page_{i + 1}.{fmt}")
            img = Image.new("RGB", (int(8.5 * dpi), int(11 * dpi)), "white")
            from PIL import ImageDraw, ImageFont

            draw = ImageDraw.Draw(img)
            try:
                font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 14)
            except OSError:
                font = ImageFont.load_default()
            y = 50
            for line in text.split("\n")[:80]:
                draw.text((50, y), line[:100], fill="black", font=font)
                y += 20
            if fmt == "jpg":
                img.save(img_path, "JPEG", quality=90)
            else:
                img.save(img_path, "PNG")
            image_paths.append(img_path)

    if len(image_paths) == 1:
        return image_paths[0]

    zip_path = get_temp_path(".zip")
    with zipfile.ZipFile(zip_path, "w") as zf:
        for p in image_paths:
            zf.write(p, os.path.basename(p))
    return zip_path


def images_to_pdf(image_paths: list[str]) -> str:
    output = get_temp_path(".pdf")
    first_img = Image.open(image_paths[0])
    if first_img.mode == "RGBA":
        first_img = first_img.convert("RGB")

    img_list = []
    for p in image_paths[1:]:
        img = Image.open(p)
        if img.mode == "RGBA":
            img = img.convert("RGB")
        img_list.append(img)

    first_img.save(output, "PDF", save_all=True, append_images=img_list, resolution=150)
    return output


def docx_to_pdf(input_path: str) -> str:
    doc = Document(input_path)
    output = get_temp_path(".pdf")
    pdf = SimpleDocTemplate(output, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    for para in doc.paragraphs:
        if para.text.strip():
            style = styles["Heading1"] if para.style.name.startswith("Heading") else styles["Normal"]
            story.append(Paragraph(para.text, style))
            story.append(Spacer(1, 6))

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells)
            story.append(Paragraph(row_text, styles["Normal"]))
        story.append(Spacer(1, 12))

    if not story:
        story.append(Paragraph("(Empty document)", styles["Normal"]))

    pdf.build(story)
    return output


def html_to_pdf(html_content: str | None = None, url: str | None = None) -> str:
    import weasyprint

    output = get_temp_path(".pdf")
    if url:
        weasyprint.HTML(url=url).write_pdf(output)
    elif html_content:
        weasyprint.HTML(string=html_content).write_pdf(output)
    return output
